"""Editable Stage 1-11 deck (pptx, 16:9).
Stage 10-11 pages: rebuilt from the deck source (deck/build.py: text boxes, rotated notes, charts as separate images,
paper background as its own picture). Other pages: rebuilt from the element specs extracted from the web deck
(edit/specs.json): text boxes, shapes, native tables, SVG illustrations as grouped native shapes, images as separate
pictures, background as its own picture. Speaker notes: notes_zh.json."""
import json, re, sys
from pathlib import Path
from lxml import etree
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE

HERE = Path(__file__).parent; DECKDIR = HERE.parent / 'deck'; OUT = Path(sys.argv[1])
specs = json.loads((HERE / 'edit/specs.json').read_text()); notes = json.loads((HERE / 'notes_zh.json').read_text()); slides_meta = json.loads((HERE / 'slides.json').read_text())
PX = 9525  # EMU per CSS px (1280 px = 13.333 in)
E = lambda v: Emu(int(round(v * PX)))
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

# ---------------- deck slides from the deck source (no save)
src = (DECKDIR / 'build.py').read_text()
src = src.replace("HERE = Path(__file__).parent", f"HERE = Path({str(DECKDIR)!r})")
src = src[:src.index("out = HERE / 'Stage10_author_readout_talk.pptx'")]
ns = {'__name__': 'deckbuild'}; exec(compile(src, 'build.py', 'exec'), ns); prs = ns['prs']
deck_ids = list(prs.slides._sldIdLst)                       # deck pages 1..27 in final order



def rr_adj(r, b):
    # hand-drawn cards use huge radii (255px) for a wobbly edge; only short boxes are real pills
    m = max(1, min(b['w'], b['h'])); a = r / m
    return min(0.5, a) if a < 0.5 or m < 48 else 0.08

def rgb(c): return RGBColor(int(c['r']), int(c['g']), int(c['b']))


def alpha(color_format, a):
    if a is None or a >= 0.999: return
    x = color_format._color._xClr
    for old in x.findall(f'{{{A_NS}}}alpha'): x.remove(old)
    el = etree.SubElement(x, f'{{{A_NS}}}alpha'); el.set('val', str(int(a * 100000)))


def set_fill(fill, c, op=1.0):
    if not c or c['a'] * op < 0.02: fill.background(); return
    fill.solid(); fill.fore_color.rgb = rgb(c); alpha(fill.fore_color, c['a'] * op)


def set_line(line, c, w_px, dash=None, op=1.0):
    if not c or not w_px or c['a'] * op < 0.02: line.fill.background(); return
    line.color.rgb = rgb(c); alpha(line.color, c['a'] * op); line.width = Pt(max(0.25, w_px * 0.75))
    if dash:
        d = str(dash)
        line.dash_style = MSO_LINE.ROUND_DOT if d in ('dotted',) or re.match(r'^\s*[0-4](px)?[ ,]+[5-9]', d) else MSO_LINE.DASH


def no_shadow(sh):
    sp = sh._element.spPr
    if sp.find(f'{{{A_NS}}}effectLst') is None: etree.SubElement(sp, f'{{{A_NS}}}effectLst')


def fill_text(tf, runs, align='left', lh=None, li=None, bullet_color=None, anchor='top', wrap=True, pad=(0, 0, 0, 0)):
    tf.word_wrap = wrap; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [E(v) for v in pad]
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if anchor == 'middle' else MSO_ANCHOR.TOP
    al = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.LEFT}.get(align, PP_ALIGN.LEFT)
    for i, para in enumerate(runs or [[]]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = al; p.space_after = Pt(0); p.space_before = Pt(0)
        if li:
            ind = E(0.9 * (para[0]['size'] if para else 20)); pPr = p._p.get_or_add_pPr(); pPr.set('marL', str(int(ind))); pPr.set('indent', str(-int(ind)))  # hanging indent: wrapped lines align with text, not bullet
        if li and i == 0:
            r = p.add_run(); r.text = '•  '; f = r.font; f.name = para[0]['font'] if para else 'Patrick Hand'; f.size = Pt((para[0]['size'] if para else 20) * 0.75)
            if bullet_color: f.color.rgb = rgb(bullet_color)
        for rr in para:
            r = p.add_run(); r.text = rr['t']; f = r.font; f.name = rr['font']; f.size = Pt(rr['size'] * 0.75); f.bold = rr['bold']; f.italic = rr['italic']
            if rr.get('color'): f.color.rgb = rgb(rr['color']); alpha(f.color, rr['color']['a'] * rr.get('op', 1))
            if rr.get('hl') and rr['hl']['a'] > 0.05:                       # highlighted span: keep a text highlight colour
                rPr = r._r.get_or_add_rPr(); hl = etree.SubElement(rPr, f'{{{A_NS}}}highlight'); c = etree.SubElement(hl, f'{{{A_NS}}}srgbClr'); c.set('val', '%02X%02X%02X' % (int(rr['hl']['r']), int(rr['hl']['g']), int(rr['hl']['b'])))
            if rr.get('ul'): f.underline = True
        if lh and para: p.line_spacing = Pt(lh * para[0]['size'] * 0.75)


def add_text_item(shapes, it):
    b = it['box']; pad = it['pad']
    single = it.get('lines', 2) <= 1
    w = b['w'] * (1.0 if it['type'] == 'shape' else 1.04)
    x = b['x'] - (0.9 * it['runs'][0][0]['size'] if it.get('li') and it['runs'] else 0)
    if it['type'] == 'shape':
        vb = it['vb']; kind = MSO_SHAPE.ROUNDED_RECTANGLE if vb['radius'] > 1 else MSO_SHAPE.RECTANGLE
        sh = shapes.add_shape(kind, E(x), E(b['y']), E(w), E(b['h']))
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = rr_adj(vb['radius'], b)
        set_fill(sh.fill, vb['fill'], vb['op']); set_line(sh.line, vb['line']['color'] if vb['line'] else None, vb['line']['w'] if vb['line'] else 0, (vb['line'] or {}).get('dash') if (vb['line'] or {}).get('dash') in ('dashed', 'dotted') else None, vb['op'])
        no_shadow(sh); sh.rotation = it['rot']
        fill_text(sh.text_frame, it['runs'], it['align'], it['lh'], it.get('li'), it.get('bulletColor'), it.get('valign', 'top'), wrap=not single, pad=pad)
        return sh
    tb = shapes.add_textbox(E(x), E(b['y'] + pad[1]), E(w + (0.9 * it['runs'][0][0]['size'] if it.get('li') else 0) - pad[0] - pad[2]), E(b['h'] - pad[1] - pad[3]))
    if pad[0]: tb.left = E(x + pad[0])
    tb.rotation = it['rot']
    fill_text(tb.text_frame, it['runs'], it['align'], it['lh'], it.get('li'), it.get('bulletColor'), wrap=not single)
    return tb


def add_svg_group(shapes, g):
    grp = shapes.add_group_shape(); grp.name = g['name'][:60]; gs = grp.shapes
    for it in g['items']:
        t = it['type']
        if t == 'rect':
            b = it['box']; kind = MSO_SHAPE.ROUNDED_RECTANGLE if it['radius'] > 0.5 else MSO_SHAPE.RECTANGLE
            sh = gs.add_shape(kind, E(b['x']), E(b['y']), E(max(b['w'], 0.5)), E(max(b['h'], 0.5)))
            if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = rr_adj(it['radius'], b)
            set_fill(sh.fill, it['fill'], it['op']); set_line(sh.line, it['stroke'], it['sw'], it['dash'], it['op']); no_shadow(sh)
        elif t == 'oval':
            b = it['box']; sh = gs.add_shape(MSO_SHAPE.OVAL, E(b['x']), E(b['y']), E(b['w']), E(b['h'])); sh.rotation = it['rot']
            set_fill(sh.fill, it['fill'], it['op']); set_line(sh.line, it['stroke'], it['sw'], it['dash'], it['op']); no_shadow(sh)
        elif t == 'line':
            c = gs.add_connector(MSO_CONNECTOR.STRAIGHT, E(it['a']['x']), E(it['a']['y']), E(it['b']['x']), E(it['b']['y']))
            set_line(c.line, it['stroke'], it['sw'], it['dash'], it['op'])
        elif t == 'poly':
            pts = [(int(round(p['x'] * PX)), int(round(p['y'] * PX))) for p in it['pts']]
            fb = gs.build_freeform(pts[0][0], pts[0][1], scale=1.0); fb.add_line_segments(pts[1:], close=it['closed']); sh = fb.convert_to_shape()
            set_fill(sh.fill, it['fill'] if it['closed'] else None, it['op']); set_line(sh.line, it['stroke'], it['sw'], it['dash'], it['op']); no_shadow(sh)
        elif t == 'svgtext':
            b = it['box']; r0 = it['runs'][0][0]; w = b['w'] * 1.15 + 6; x = b['x'] - (w - b['w']) * (0.5 if it['anchor'] == 'middle' else (1 if it['anchor'] == 'end' else 0))
            tb = gs.add_textbox(E(x), E(b['y']), E(w), E(b['h'])); tb.rotation = it['rot']
            fill_text(tb.text_frame, it['runs'], {'middle': 'center', 'end': 'right'}.get(it['anchor'], 'left'), wrap=False)
    return grp


def add_table(shapes, tspec):
    rows = tspec['rows']; nr = len(rows); nc = max(sum(c['span'] for c in r) for r in rows)
    b = tspec['box']; gf = shapes.add_table(nr, nc, E(b['x']), E(b['y']), E(b['w']), E(b['h'])); tbl = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    for a in ('firstRow', 'bandRow', 'firstCol', 'lastRow', 'lastCol', 'bandCol'): tblPr.set(a, '0')
    sid = tblPr.find(f'{{{A_NS}}}tableStyleId')
    if sid is not None: sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'              # "No Style, No Grid"
    else: etree.SubElement(tblPr, f'{{{A_NS}}}tableStyleId').text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    widest = max(rows, key=lambda r: len(r))
    for j, c in enumerate(widest[:nc]): tbl.columns[j].width = E(c['box']['w'])
    for i, r in enumerate(rows):
        tbl.rows[i].height = E(max(c['box']['h'] for c in r)); j = 0
        for c in r:
            cell = tbl.cell(i, j)
            if c['span'] > 1: cell.merge(tbl.cell(i, j + c['span'] - 1))
            fill_text(cell.text_frame, c['runs'], c['align'], None, pad=[min(c['pad'][0], 2), c['pad'][1], min(c['pad'][2], 2), c['pad'][3]])  # LibreOffice wraps at exact widths
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE if c['valign'] == 'middle' else MSO_ANCHOR.TOP
            cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = E(min(c['pad'][0], 2)), E(c['pad'][1]), E(min(c['pad'][2], 2)), E(c['pad'][3])  # cells ignore bodyPr insets; default is 0.1 in
            tcPr = cell._tc.get_or_add_tcPr()
            for k in list(tcPr): tcPr.remove(k)
            for side in ('lnL', 'lnR', 'lnT', 'lnB'):
                ln = etree.SubElement(tcPr, f'{{{A_NS}}}{side}')
                bb = c['bb']
                if side == 'lnB' and bb['w'] > 0 and bb['color']:
                    ln.set('w', str(int(bb['w'] * 0.75 * 12700))); sf = etree.SubElement(ln, f'{{{A_NS}}}solidFill'); cl = etree.SubElement(sf, f'{{{A_NS}}}srgbClr')
                    cl.set('val', '%02X%02X%02X' % (int(bb['color']['r']), int(bb['color']['g']), int(bb['color']['b'])))
                    if bb['color']['a'] < 0.999: etree.SubElement(cl, f'{{{A_NS}}}alpha').set('val', str(int(bb['color']['a'] * 100000)))
                    if bb['style'] in ('dashed', 'dotted'): etree.SubElement(ln, f'{{{A_NS}}}prstDash').set('val', 'dash' if bb['style'] == 'dashed' else 'sysDot')
                else:
                    ln.set('w', '0'); etree.SubElement(ln, f'{{{A_NS}}}noFill')
            if c['bg'] and c['bg']['a'] > 0.02:
                sf = etree.SubElement(tcPr, f'{{{A_NS}}}solidFill'); cl = etree.SubElement(sf, f'{{{A_NS}}}srgbClr'); cl.set('val', '%02X%02X%02X' % (int(c['bg']['r']), int(c['bg']['g']), int(c['bg']['b'])))
            else: etree.SubElement(tcPr, f'{{{A_NS}}}noFill')
            j += c['span']
    return gf


BLANK = prs.slide_layouts[6]
new_ids = {}
for idx_s, spec in specs.items():
    s = prs.slides.add_slide(BLANK)
    bgp = s.shapes.add_picture(spec['bg'], 0, 0, prs.slide_width, prs.slide_height); bgp.name = 'background'
    texts = [x for x in spec['items'] if x['type'] == 'text']
    for it in spec['items']:
        if it['type'] == 'shape' and it.get('runs'):
            # card text comes after child headings that are drawn as their own text boxes: start below them
            b = it['box']; kids = [x['box'] for x in texts if b['x'] - 1 <= x['box']['x'] and x['box']['x'] + x['box']['w'] <= b['x'] + b['w'] + 1 and b['y'] - 1 <= x['box']['y'] < b['y'] + b['h'] / 2]
            if kids: it['pad'] = [it['pad'][0], max(it['pad'][1], max(k['y'] + k['h'] for k in kids) - b['y'] + 2), it['pad'][2], it['pad'][3]]
        t = it['type']
        if t in ('text', 'shape'): add_text_item(s.shapes, it)
        elif t == 'group': add_svg_group(s.shapes, it)
        elif t == 'table': add_table(s.shapes, it)
        elif t == 'img':
            b = it['box']; pic = s.shapes.add_picture(it['src'], E(b['x']), E(b['y']), E(b['w']), E(b['h'])); pic.rotation = it['rot']; pic.name = 'chart image'
    new_ids[int(idx_s)] = s
all_ids = list(prs.slides._sldIdLst)
id_of_new = {k: all_ids[len(deck_ids) + n] for n, k in enumerate(sorted(new_ids, key=lambda k: list(specs).index(str(k))))}
order = []
for m in slides_meta:
    order.append(deck_ids[m['deck'] - 1] if m['kind'] == 'deck' else id_of_new[m['i']])
lst = prs.slides._sldIdLst
for e in list(lst): lst.remove(e)
used = set(id(e) for e in order)
for e in order: lst.append(e)
for e in all_ids:                                            # deck pages not used in the merged deck: drop them
    if id(e) not in used: prs.part.drop_rel(e.rId)
# notes and the two re-checked pages' stage stamps
for k, sl in enumerate(prs.slides, 1):
    nt = notes[str(k)]; m = slides_meta[k - 1]
    src_txt = f"deck 第 {m['deck']} 页（按源文件重建）" if m['kind'] == 'deck' else '合并 PPT 自身页面（按网页元素重建）'
    sl.notes_slide.notes_text_frame.text = f"第 {k} / {len(slides_meta)} 页 · {src_txt} · " + ('【新写讲稿】' if nt['new'] else '按 deck 讲稿备注翻译（已按复核结果更新）') + '\n\n' + nt['text']
prs.save(OUT); print('saved', OUT, len(prs.slides), 'slides')
